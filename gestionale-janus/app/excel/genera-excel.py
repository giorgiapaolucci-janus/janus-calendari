#!/usr/bin/env python3
"""Uso: python3 genera-excel.py <cartella-dati> <stato.json> <uscita.xlsx>
cartella-dati contiene mattina.json sera.json docenti.json (generati da genera-studenti.js)"""
import json, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
dd, stf, out = sys.argv[1:4]
J = lambda p: json.load(open(p, encoding='utf8'))
doc = J(f'{dd}/docenti.json'); mat = J(f'{dd}/mattina.json'); ser = J(f'{dd}/sera.json')
sw = J('software.json'); st = J(stf)
NOMI = doc['nomi']; GG = ['', 'Lunedì', 'Martedì', 'Mercoledì', 'Giovedì', 'Venerdì']
zoom = {}
for f in (mat, ser):
    for l in f['lezioni']:
        r = l['room'].replace('AULA ', '')
        if r.isdigit(): zoom.setdefault(int(r), l['zoomUrl'])
HF = PatternFill('solid', fgColor='2B2A4C'); HFont = Font(bold=True, color='FFFFFF')
thin = Side(style='thin', color='D0D0DA'); BR = Border(left=thin, right=thin, top=thin, bottom=thin)
def sheet(wb, nome, intest, righe, larg, link_col=None):
    ws = wb.create_sheet(nome)
    ws.append(intest)
    for c in ws[1]: c.fill = HF; c.font = HFont; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = BR
    for r in righe: ws.append(r)
    for i, w in enumerate(larg, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BR; c.alignment = Alignment(vertical='top', wrap_text=True)
            if link_col and c.column == link_col and c.value and str(c.value).startswith('http'):
                c.hyperlink = c.value; c.font = Font(color='1F4FBF', underline='single')
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    return ws
def aula_txt(r):
    return 'Aula docenti' if r['esercizi'] else f"Aula {r['aula']}"
def righe(rows):
    rows = sorted(rows, key=lambda r: (r['day'], r['hour'], r['anno'], r['doc']))
    return [[GG[r['day']], f"{r['hour']:02d}:00 – {r['hour']:02d}:50", r['anno'] + ' anno', r['ind'], r['materia'], NOMI.get(r['doc'], r['doc']),
             aula_txt(r), 'Esercizi' if r['esercizi'] else (r['note'] or ''), zoom.get(r['aula'], ''), r['id']] for r in rows]
I1 = ['Giorno', 'Orario', 'Anno', 'Indirizzi', 'Materia', 'Docente', 'Aula', 'Note', 'Link lezione (Zoom)', 'ID']
L1 = [11, 14, 10, 24, 24, 20, 13, 14, 60, 7]
wb = Workbook(); wb.remove(wb.active)
# Info
ws = wb.create_sheet('Info')
agg = datetime.datetime.now().strftime('%d/%m/%Y alle %H:%M')
info = [['CALENDARIO LEZIONI UFFICIALE – Istituto Janus'], [f'Aggiornato il {agg}'], [''],
 ['Questo file viene generato dal software dei calendari ed è aggiornato a ogni modifica. Non modificarlo a mano: le modifiche si fanno nel software.'],
 [''], ['Schede'], ['Mattina – tutte le lezioni della mattina (9:00–13:00)'], ['Sera – tutte le lezioni della sera (17:00–21:00)'],
 ['Orari docenti – settimana tipo di ogni docente'], ['Aule e link – link Zoom di ogni aula'], ['Turni docenti – presenza / smart working'],
 ['Modifiche e sostituzioni – assenze, sostituzioni e cambi registrati'], [''], ['Legenda turni: P = presenza, S = smart working, - = nessun turno'],
 ['Le lezioni con "Aula docenti" sono in modalità Esercizi: il docente è in un\'altra aula, chi è in sede svolge gli esercizi in Aula docenti.']]
for r in info: ws.append(r)
ws['A1'].font = Font(bold=True, size=16); ws['A6'].font = Font(bold=True); ws.column_dimensions['A'].width = 120
for r in ws.iter_rows():
    for c in r: c.alignment = Alignment(wrap_text=True, vertical='top')
sheet(wb, 'Mattina', I1, righe(doc['diurno']), L1, 9)
sheet(wb, 'Sera', I1, righe(doc['serale']), L1, 9)
# Orari docenti
tutte = [dict(r, fascia='M') for r in doc['diurno']] + [dict(r, fascia='S') for r in doc['serale']]
rr = []
for sigla, nome in sorted(NOMI.items(), key=lambda x: x[1]):
    mie = [r for r in tutte if r['doc'] == sigla]
    for h in sorted({r['hour'] for r in mie}):
        cella = lambda d: '\n'.join(f"{r['materia']} · {r['anno']} · {r['ind']} · {aula_txt(r)}" for r in mie if r['day'] == d and r['hour'] == h)
        rr.append([nome, f'{h:02d}:00 – {h:02d}:50'] + [cella(d) for d in range(1, 6)])
ws = sheet(wb, 'Orari docenti', ['Docente', 'Orario'] + GG[1:], rr, [22, 14, 30, 30, 30, 30, 30])
# Aule
ar = []
for n in sorted(zoom):
    docs = sorted({NOMI.get(r['doc'], r['doc']) for r in doc['diurno'] + doc['serale'] if r['aula'] == n and not r['esercizi']})
    ar.append([f'Aula {n}', zoom[n], ', '.join(docs)])
ar.append(['Aula docenti', 'Nessun link: è la sala in sede per chi svolge gli esercizi; il docente è collegato nella sua aula', ''])
sheet(wb, 'Aule e link', ['Aula', 'Link Zoom', 'Docenti che la usano'], ar, [16, 75, 70], 2)
# Turni
tr = []
for d in sw['DOCENTI']:
    m = sw['TURNI']['m'].get(d['id'], '-----')
    tr.append([d['nome'], 'Sì' if 'M' in d['cal'] else 'No', 'Sì' if d['sera'] else 'No'] + list(m))
sheet(wb, 'Turni docenti', ['Docente', 'Mattina', 'Sera'] + [g for g in GG[1:]], tr, [22, 10, 8, 11, 11, 11, 11, 11])
# Modifiche
mr = []; byid = {d['id']: d['nome'] for d in sw['DOCENTI']}
for a in st.get('assenze', []):
    mr.append([a.get('dal'), a.get('al'), 'Assenza', byid.get(a['docenteId'], a['docenteId']), a.get('tipo', ''), a.get('nota', '')])
for s in st.get('sost', []):
    for lid, did in (s.get('map') or {}).items():
        mr.append([s.get('data'), s.get('data'), 'Sostituzione', byid.get(did, 'nessun sostituto') if did else 'nessun sostituto', lid, ''])
for m in st.get('mods', []):
    mr.append([(m.get('creata') or '')[:10], '', 'Modifica permanente', '', m.get('lessonId', ''), m.get('motivo', '')])
mr.sort(key=lambda r: str(r[0]))
sheet(wb, 'Modifiche e sostituzioni', ['Dal', 'Al', 'Tipo', 'Docente', 'Dettaglio', 'Note'], mr, [12, 12, 20, 24, 20, 80])
wb.save(out); print('ok', out, {ws.title: ws.max_row - 1 for ws in wb if ws.title != 'Info'})
